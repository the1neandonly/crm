import sys
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from brand import *
OUT = sys.argv[1]
INK, VIO = RGBColor(0x24, 0x24, 0x2B), RGBColor(0x63, 0x37, 0xFF)
d = Document(); sec = d.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(20); sec.top_margin = Mm(50); sec.bottom_margin = Mm(30); sec.header_distance = Mm(16); sec.footer_distance = Mm(12)
st = d.styles['Normal']; st.font.name = 'Manrope'; st.font.size = Pt(10.5); st.font.color.rgb = RGBColor(0x3a, 0x3a, 0x44)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Manrope')
for lvl, sz in (('Heading 1', 24), ('Heading 2', 15)):
    h = d.styles[lvl]; h.font.name = 'Syne'; h.font.bold = True; h.font.size = Pt(sz); h.font.color.rgb = INK
# header: logo left, contact right
hdr = sec.header; t = hdr.add_table(rows=1, cols=2, width=Mm(170)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
c0, c1 = t.rows[0].cells
c0.paragraphs[0].add_run().add_picture(f'{OUT}/logo/png/horizontal/bayna-horizontal-color-800px.png', width=Mm(54))
p = c1.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run('Bayna Agency\n'); r.bold = True; r.font.size = Pt(7.6); r.font.color.rgb = INK
r = p.add_run(f'{CITY}\n{PHONE}\n{MAIL} · {WEB}'); r.font.size = Pt(7.6); r.font.color.rgb = RGBColor(0x55, 0x55, 0x5F)
hdr.paragraphs[0].text = ''
# footer
fp = sec.footer.paragraphs[0]
r = fp.add_run('Bayna Agency'); r.bold = True; r.font.size = Pt(7); r.font.color.rgb = VIO
r = fp.add_run(' — Marketing & Communication · باينة      RC 00/00-0000000 · NIF 000000000000000'); r.font.size = Pt(7); r.font.color.rgb = RGBColor(0x8a, 0x8a, 0x96)
# body sample
d.add_paragraph('Oran, le __ / __ / ____').runs[0].font.color.rgb = RGBColor(0x8a, 0x8a, 0x96)
d.add_paragraph('')
d.add_paragraph('Objet : ').runs[0].bold = True
d.add_paragraph('Madame, Monsieur,')
d.add_paragraph('Rédigez votre courrier ici. Le titre utilise Syne Bold, le texte Manrope Regular 10,5 pt. Installez les polices du dossier /fonts avant d’ouvrir ce modèle.')
d.add_paragraph('Cordialement,')
d.add_paragraph(f'{NAME}\n{ROLE}')
d.save(f'{OUT}/print/letterhead/Bayna-letterhead-A4-EDITABLE.docx'); print('docx ok')
